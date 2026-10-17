from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];p=json.loads((root/'data/report_inputs.json').read_text());k=p['kpis']
def document(title):
 d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  style=d.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(15)
 
 for elem in [d.styles.element,d.element]:
  for border in elem.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
 d.add_heading(title,0);return d

def table(d,head,rows):
 t=d.add_table(rows=1,cols=len(head));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,head):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 return t

d=document('Institutional partnership proposal')
d.add_paragraph('CampusBridge Demo to Simulated Ahmedabad University 426\nPrepared by Harsh Vardhan Tomar | Portfolio simulation | 30 September 2026')
d.add_paragraph('We propose a 12-month institutional discovery and enquiry-management campaign to help your admissions team understand prospective-student interest and follow up consistently. This sample proposal relates to opportunity O0426. It is a fictional commercial discussion, not an offer to a real university.')
d.add_heading('Proposed scope',1)
for text in ['Institution profile and programme pages with admission requirements, approved content and enquiry routing.','Two admissions webinars per quarter, supported by institution-approved campaign materials.','Monthly reporting on consented enquiries, response times, qualified applicants and application completion.','Weekly lead-quality review with an assigned admissions coordinator and a shared issue log.']:d.add_paragraph(text,style='List Bullet')
d.add_heading('Commercial assumptions',1)
table(d,['Item','Proposed basis'],[['Annual service fee','INR 1,223,000 excluding applicable taxes'],['Term','365 days beginning on agreed signature and service activation'],['Invoicing assumption','25% at activation and 25% at the start of each subsequent quarter'],['Paid media','Excluded; separate written approval and budget required'],['Sales status','Open opportunity; the proposed fee is not booked revenue']])
d.add_heading('Potential benefit and measurement',1)
d.add_paragraph('The programme is intended to improve visibility into the enquiry funnel and make admissions follow-up measurable. No enrolment uplift or revenue return is promised. The first 30 days establish a baseline before the parties approve campaign-volume goals.')
d.add_page_break();d.add_heading('Delivery and success criteria',0)
table(d,['Period','Deliverable','Acceptance measure'],[['Days 1–30','Profile, routing and consent review','University approves content and tests enquiry delivery'],['Months 2–3','Campaign launch and baseline review','Weekly reporting complete; duplicate and invalid leads classified'],['Months 4–12','Campaign optimisation and quarterly review','Agreed activities delivered and corrective actions tracked']])
d.add_heading('Proposed service measures',1)
d.add_paragraph('Operational targets to negotiate: route 95% of valid enquiries within one working day, achieve complete source tagging for 98% of routed enquiries, and deliver one management report each month. These are proposed targets, not historical performance. Enrolment conversion is reported as verified enrolments divided by eligible enquiries, with the agreed attribution window shown.')
d.add_heading('Responsibilities and safeguards',1)
d.add_paragraph('The university approves content, supplies accurate programme information and appoints an admissions owner. CampusBridge Demo maintains campaign reporting and records consent and source information. A real agreement would specify data access, retention, grievance handling, service scope, taxes and termination terms. Student contact data is absent from this portfolio dataset.')
d.add_heading('Decision requested',1)
d.add_paragraph('Arrange a discovery meeting to confirm programme priorities, available capacity and a baseline measurement plan. Agree scope and pricing only after procurement review. This proposal does not change the opportunity to Closed Won.')
d.save(root/'reports/Institutional_Partnership_Proposal.docx')
d=document('Sales performance business review')
d.add_paragraph('CampusBridge Demo | Management review | 01 October 2025 to 30 September 2026\nPrepared by Harsh Vardhan Tomar | All commercial records are simulated')
d.add_paragraph('Prioritise pipeline hygiene before expanding outreach. The simulated team has 370 open opportunities, but 270 have spent more than 30 days in their current stage and 222 have overdue next actions. A large gross pipeline therefore does not justify treating every deal as near-term revenue.')
table(d,['Measure','Snapshot result'],[['Institution prospects / activities','600 / 4,399'],['Active pipeline','INR 109,230,000'],['Signed partnership value','INR 16,886,000 across 57 wins'],['Recognized service revenue','INR 5,736,156.17'],['Closed-deal win rate','57 / 230 = 24.78%'],['Average won deal / mean won cycle','INR 296,245.61 / 124.93 days']])
d.add_heading('Conversion bottleneck',1)
d.add_paragraph('Proposal Sent to Negotiation has the lowest observed adjacent-stage conversion among stages with at least 20 entries: 75 of 146 opportunities, or 51.37%. Review scope clarity, procurement objections and the next decision date. Recent opportunities remain unfinished, so this ratio is not a final probability of progression.')
d.add_heading('Territory allocation',1)
d.add_paragraph('North holds INR 29,212,000 of active pipeline and 57 overdue actions. South has 16 wins out of 47 closed deals (34.04%), compared with East at 5 of 38 (13.16%). These differences are descriptive simulation results. Lead mix and deal maturity must be considered before reallocating staff. Clear North’s backlog and examine East’s loss reasons before a new territory pilot.')
d.add_page_break();d.add_heading('Forecast and management actions',0)
table(d,['Q4 2026 scenario','New signed value INR'],[[x['scenario'],f"{x['forecast']:,.0f}"] for x in p['forecast']])
d.add_paragraph('The forecast covers 108 open opportunities expected to close in October–December 2026. It excludes 224 past-due deals carrying INR 71,044,000 and deals scheduled after the quarter. Stage probabilities are assumptions: 3%, 10%, 18%, 30%, 50% and 70%. Conservative and optimistic scenarios scale them by 0.7 and 1.3, capped at 100%. These are bookings scenarios, not confidence intervals or recognized-revenue estimates.')
# Value is sourced below to guard copy consistency.
stale=sum(x['value'] for x in p['analytics'] if x['status']=='Open' and x['expected_close']<='2026-09-30')
for para in d.paragraphs:
 if '71,044,000' in para.text:para.text=para.text.replace('71,044,000',f'{stale:,.0f}')
table(d,['Owner','Next review action','Metric to monitor'],[['Sales representatives','Validate overdue tasks within 7 days','222 overdue opportunities'],['Sales manager','Requalify stalled deals at weekly review','270 opportunities >30 days in stage'],['Commercial lead','Audit proposal objections','75 / 146 advanced to Negotiation'],['Sales operations','Refresh expected close dates before forecast','224 past-due open deals'],['Territory lead','Review account O0426 and confirm next step','80.21 priority score; INR 1,223,000 potential']])
score=next(x['score'] for x in p['analytics'] if x['opportunity_id']=='O0426')
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   if '80.21' in c.text:c.text=c.text.replace('80.21',str(score))
d.add_heading('Evidence and limits',1)
d.add_paragraph('Source: data/partnerships.sqlite and data/summary.json, deterministic seed 20261017. SQL and Python reconcile signed value and accrual. The dataset has one opportunity per institution, forward-only stages and constant opportunity values. No causal improvement, real client result, market-size estimate or model accuracy is claimed. Full formulas and denominators appear in docs/KPI_DEFINITIONS.md.')
d.save(root/'reports/Sales_Performance_Review.docx')
